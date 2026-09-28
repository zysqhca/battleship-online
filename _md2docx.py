# -*- coding: utf-8 -*-
import re
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

def set_cn_font(run, size=None, bold=None, color=None):
    run.font.name = '微软雅黑'
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)

def is_separator(cells):
    return all(re.fullmatch(r'[\s:\-|]+', c) for c in cells)

def parse_md(path):
    with open(path, encoding='utf-8') as f:
        lines = f.read().split('\n')
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        if line.strip() == '' or line.strip() == '---':
            i += 1
            continue
        if line.startswith('### '):
            blocks.append(('h3', line[4:].strip()))
        elif line.startswith('## '):
            blocks.append(('h2', line[3:].strip()))
        elif line.startswith('# '):
            blocks.append(('h1', line[2:].strip()))
        elif line.startswith('- '):
            items = []
            while i < n and lines[i].lstrip().startswith('- '):
                items.append(lines[i].lstrip()[2:].strip())
                i += 1
            blocks.append(('ul', items))
            continue
        elif line.startswith('|'):
            rows = []
            while i < n and lines[i].lstrip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(('table', rows))
            continue
        elif line.startswith('> '):
            blocks.append(('quote', line[2:].strip()))
        else:
            blocks.append(('p', line))
        i += 1
    return blocks

def add_runs(p, text, size=11, bold=False, color=None):
    # 处理 **加粗**
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            set_cn_font(run, size=size, bold=True, color=color)
        else:
            run = p.add_run(part)
            set_cn_font(run, size=size, bold=bold, color=color)

def build(src, dst):
    blocks = parse_md(src)
    d = docx.Document()
    for kind, content in blocks:
        if kind == 'h1':
            p = d.add_paragraph()
            add_runs(p, content, size=20, bold=True, color=(31, 41, 55))
        elif kind == 'h2':
            p = d.add_paragraph()
            add_runs(p, content, size=15, bold=True, color=(99, 102, 241))
        elif kind == 'h3':
            p = d.add_paragraph()
            add_runs(p, content, size=12, bold=True, color=(31, 41, 55))
        elif kind == 'p':
            p = d.add_paragraph()
            add_runs(p, content, size=11)
        elif kind == 'ul':
            for it in content:
                p = d.add_paragraph(style='List Bullet')
                add_runs(p, it, size=11)
        elif kind == 'quote':
            p = d.add_paragraph()
            add_runs(p, content, size=10, color=(107, 114, 128))
        elif kind == 'table':
            rows = []
            for raw in content:
                cells = [c.strip() for c in raw.strip().strip('|').split('|')]
                if is_separator(cells):
                    continue
                rows.append(cells)
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            tbl = d.add_table(rows=len(rows), cols=ncol)
            tbl.style = 'Light Grid Accent 1'
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.rows[ri].cells[ci]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    txt = r[ci] if ci < len(r) else ''
                    add_runs(p, txt, size=10, bold=(ri == 0))
            d.add_paragraph()
    d.save(dst)

if __name__ == '__main__':
    build('海战棋完整规则.md', '海战棋完整规则.docx')
    print('ok')
